"""Uploaded files: celebrity sheets and client profiles."""
import io
import zipfile

import openpyxl
import pytest

from app.data import repositories as repo
from app.services import celebrities as celebs
from app.services import documents, files


def _xlsx(sheets: dict) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_sheet_with_arabic_headers_and_a_title_row():
    data = _xlsx({"مشاهير": [["قائمة المشاهير المعتمدة"], [],
                              ["م", "اسم المشهور", "المنصة", "عدد المتابعين", "المجال", "المدينة", "السعر"],
                              [1, "سارة", "سناب شات", 1200000, "لايف ستايل", "الرياض", 9000],
                              [2, "فهد", "تيك توك", 800000, "سيارات", "جدة", 5000]]})
    people = files.parse_celebrity_sheet("list.xlsx", data)
    assert [p["name"] for p in people] == ["سارة", "فهد"]
    assert people[0]["followers"] == "1,200,000" and people[0]["city"] == "الرياض"
    assert people[0]["extra"]["السعر"] == "9,000"  # kept, but never sent to the writer (see author._candidates_table)


def test_sheet_with_one_column_per_platform_and_several_tabs():
    data = _xlsx({"الرياض": [["الاسم", "سناب", "انستقرام"], ["نورة", "@noura", "450K"]],
                  "جدة": [["Name", "TikTok followers"], ["Ali", "2M"]]})
    people = files.parse_celebrity_sheet("multi.xlsx", data)
    noura = next(p for p in people if p["name"] == "نورة")
    assert noura["handle"] == "@noura" and "إنستقرام" in noura["platforms"] and noura["category"] == "الرياض"
    ali = next(p for p in people if p["name"] == "Ali")
    assert "تيك توك" in ali["platforms"] and ali["category"] == "جدة"


def test_csv_and_import_dedupes():
    repo.celebrities.clear()
    csv_bytes = "الاسم,المنصة,المتابعين\nسارة,سناب,1M\nسارة,سناب,1M\nفهد,تيك توك,800K\n".encode("utf-8-sig")
    res = celebs.import_file("list.csv", csv_bytes)
    assert res["added"] == 2 and res["skipped"] == 1 and res["total"] == 2
    assert celebs.import_file("again.csv", csv_bytes)["added"] == 0
    repo.celebrities.clear()


def test_bad_sheets_explain_themselves():
    with pytest.raises(files.FileError, match="xlsx"):
        files.parse_celebrity_sheet("old.xls", b"x")
    with pytest.raises(files.FileError, match="عمودًا للاسم"):
        files.parse_celebrity_sheet("empty.csv", b"\n\n")


def test_profile_from_word_and_powerpoint():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", '<w:document><w:body><w:p><w:r><w:t>شركة الدروب</w:t></w:r></w:p>'
                                         '<w:p><w:r><w:t>نؤجر السيارات في الرياض &amp; جدة</w:t></w:r></w:p></w:body></w:document>')
    info = files.extract_document("profile.docx", buf.getvalue())
    assert info["text"] == "شركة الدروب\nنؤجر السيارات في الرياض & جدة"

    from pptx import Presentation
    prs = Presentation()
    for t in ("من نحن", "خدماتنا"):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = t
    out = io.BytesIO()
    prs.save(out)
    info = files.extract_document("deck.pptx", out.getvalue())
    assert info["pages"] == 2 and "من نحن" in info["text"] and "خدماتنا" in info["text"]


def test_store_profile_and_reject_unsupported():
    doc = documents.store("about.txt", "نبذة عن العميل: متجر قهوة مختصة في الرياض.".encode("utf-8"), keep_file=False)
    assert doc["chars"] > 10 and repo.documents.get(doc["id"])["text"].startswith("نبذة")
    with pytest.raises(files.FileError):
        documents.store("photo.png", b"\x89PNG")
    with pytest.raises(files.FileError, match="لا يحتوي على نص"):
        documents.store("blank.txt", b"   ")
