"""PowerPoint export.

Builds an editable 16:9 deck (native text boxes, shapes and tables — no screenshots) in AHDAF's
visual language: purple cover/dividers/closing, white content slides, lilac cards, purple table
headers. Arabic paragraphs are written right-to-left with the brand font set for complex scripts.
A small layout engine estimates text heights and splits long lists/tables across slides.
"""
from __future__ import annotations

import io
import math
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from app import config

BRAND = config.WEB_DIR / "assets" / "brand"
UPLOADS = config.UPLOAD_DIR

PURPLE, DEEP, NIGHT, MID = "4C0192", "2B0560", "1A033F", "6A1FB0"
ACCENT, LILAC, LILAC2, LINE = "B77BFF", "F4F0FA", "EAE1F8", "E6E0EF"
INK, INK2, MUTED, WHITE = "1E1630", "3B3350", "6E6585", "FFFFFF"
HL = "FFF1B8"
LVL = {"hi": ("FBE6E3", "C0392B"), "md": ("FDF3DC", "9A6412"), "lo": ("E6F4EC", "2F855A")}

SW, SH = 13.333, 7.5
X0, CW = 0.7, 13.333 - 1.4
Y0, YMAX = 1.62, 6.82
GAP = 0.18
_PH = re.compile(r"(\[\[[^\]]+\]\])")
_LATIN = re.compile(r"([0-9A-Za-z][0-9A-Za-z+.,%:/\-–()@_&' ]*[0-9A-Za-z%)]|[0-9A-Za-z])")


class Deck:
    def __init__(self, font: str):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SW)
        self.prs.slide_height = Inches(SH)
        self.blank = self.prs.slide_layouts[6]
        self.font = font or "Saudi Serif VF"
        self.page = 0

    # ------------------------------------------------------------------ primitives
    def slide(self, bg: str = WHITE):
        s = self.prs.slides.add_slide(self.blank)
        fill = s.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(bg)
        self.page += 1
        return s

    def rect(self, s, x, y, w, h, fill, radius=0.12, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None, alpha=None):
        shp = s.shapes.add_shape(shape if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        if radius and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            shp.adjustments[0] = min(0.5, radius / max(0.01, min(w, h)))
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill)
        if alpha is not None:
            clr = shp._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            a = etree.SubElement(clr, qn("a:alpha"))
            a.set("val", str(int(alpha * 100000)))
        if line:
            shp.line.color.rgb = RGBColor.from_string(line)
            shp.line.width = Pt(0.75)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        shp.text_frame.text = ""
        return shp

    def pic(self, s, path, x, y, w=None, h=None, alpha=None):
        p = Path(path)
        if not p.exists():
            return None
        kw = {}
        if w:
            kw["width"] = Inches(w)
        if h:
            kw["height"] = Inches(h)
        pic = s.shapes.add_picture(str(p), Inches(x), Inches(y), **kw)
        if alpha is not None:
            blip = pic._element.find(".//" + qn("a:blip"))
            amf = etree.SubElement(blip, qn("a:alphaModFix"))
            amf.set("amt", str(int(alpha * 100000)))
        return pic

    def text(self, s, x, y, w, h, paras, anchor="top", margin=0.04, autofit=False):
        """paras: list of dicts {text, size, bold, color, align, bullet, space_after, line}"""
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, Inches(margin))
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        body = tf._txBody.find(qn("a:bodyPr"))
        for child in list(body):
            body.remove(child)
        if autofit:
            etree.SubElement(body, qn("a:normAutofit"))
        for i, pd in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            self._para(p, pd)
        return tb

    def _para(self, p, pd):
        pPr = p._p.get_or_add_pPr()
        align = pd.get("align", "r")
        pPr.set("algn", align)
        pPr.set("rtl", "0" if pd.get("ltr") else "1")
        if pd.get("bullet"):
            pPr.set("marL", str(Emu(Inches(0.22))))
            pPr.set("indent", str(-Emu(Inches(0.22))))
        p.line_spacing = pd.get("line", 1.2)
        p.space_after = Pt(pd.get("space_after", 4))
        if pd.get("bullet"):
            bc = etree.SubElement(pPr, qn("a:buClr"))
            etree.SubElement(bc, qn("a:srgbClr")).set("val", ACCENT)
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
        for part in _PH.split(str(pd.get("text") or "")):
            if not part:
                continue
            is_ph = part.startswith("[[") and part.endswith("]]")
            chunk = part[2:-2] if is_ph else part
            for k, seg in enumerate(_LATIN.split(chunk)):
                if seg:
                    self._run(p, seg, pd.get("size", 14), pd.get("bold", False), pd.get("color", INK),
                              highlight=HL if is_ph else None, latin=bool(k % 2))

    def _run(self, p, text, size, bold, color, highlight=None, latin=False):
        r = p.add_run()
        r.text = text
        rPr = r._r.get_or_add_rPr()
        for c in list(rPr):
            rPr.remove(c)
        rPr.set("lang", "en-US" if latin else "ar-SA")
        rPr.set("altLang", "ar-SA" if latin else "en-US")
        rPr.set("sz", str(int(size * 100)))
        rPr.set("b", "1" if bold else "0")
        sf = etree.SubElement(rPr, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
        if highlight:
            hl = etree.SubElement(rPr, qn("a:highlight"))
            etree.SubElement(hl, qn("a:srgbClr")).set("val", highlight)
        etree.SubElement(rPr, qn("a:latin")).set("typeface", self.font)
        etree.SubElement(rPr, qn("a:cs")).set("typeface", self.font)

    def table(self, s, x, y, w, col_w, header, rows, row_h, size=11.5, first_bold=True, cell_fmt=None):
        n_rows, n_cols = len(rows) + 1, len(header)
        total_h = 0.42 + sum(row_h)
        gt = s.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(total_h))
        tbl = gt.table
        tblPr = tbl._tbl.tblPr
        tblPr.set("rtl", "1")
        tblPr.set("firstRow", "1")
        tblPr.set("bandRow", "0")
        style = tblPr.find(qn("a:tableStyleId"))
        if style is not None:
            style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid"
        for ci, cw in enumerate(col_w):
            tbl.columns[ci].width = Inches(cw)
        tbl.rows[0].height = Inches(0.42)
        for ri, rh in enumerate(row_h, start=1):
            tbl.rows[ri].height = Inches(rh)
        for ci, label in enumerate(header):
            self._cell(tbl.cell(0, ci), label, size - 0.5, True, WHITE, PURPLE)
        for ri, row in enumerate(rows, start=1):
            for ci, val in enumerate(row):
                fmt = (cell_fmt or {}).get(ci, {})
                bg = fmt.get("bg") or ("FBF9FE" if ri % 2 == 0 else WHITE)
                color = fmt.get("color") or (DEEP if ci == 0 and first_bold else INK)
                self._cell(tbl.cell(ri, ci), val, size, ci == 0 and first_bold or fmt.get("bold", False), color, bg,
                           align=fmt.get("align", "r"))
        return total_h

    def _cell(self, cell, text, size, bold, color, bg, align="r"):
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(bg)
        cell.margin_left = cell.margin_right = Inches(0.08)
        cell.margin_top = cell.margin_bottom = Inches(0.05)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        self._para(p, {"text": str(text if text is not None else ""), "size": size, "bold": bold, "color": color, "align": align, "space_after": 0, "line": 1.1})
        tcPr = cell._tc.get_or_add_tcPr()
        for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
            ln = tcPr.find(qn(edge))
            if ln is None:
                ln = etree.SubElement(tcPr, qn(edge))
            ln.set("w", str(Emu(Pt(0.5))))
            for c in list(ln):
                ln.remove(c)
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr")).set("val", LINE if edge in ("a:lnB", "a:lnT") else bg)
        # tcPr children must keep schema order: lnL, lnR, lnT, lnB, ..., fill
        fills = [c for c in tcPr if c.tag in (qn("a:solidFill"), qn("a:noFill"))]
        for f in fills:
            tcPr.remove(f)
            tcPr.append(f)


# ---------------------------------------------------------------------- text metrics
def _plain(text) -> str:
    return re.sub(r"\[\[|\]\]", "", str(text or ""))


def lines(text, width_in, pt) -> int:
    cpl = max(8, int(width_in * 72 / (pt * 0.47)))
    total = 0
    for para in _plain(text).split("\n"):
        total += max(1, math.ceil(len(para) / cpl))
    return total


def th(text, width_in, pt, lh=1.38) -> float:
    if not str(text or "").strip():
        return 0.0
    return lines(text, width_in, pt) * pt * lh / 72 + 0.06


# ---------------------------------------------------------------------- slide chrome
def content_slide(d: Deck, section: dict, num: str, first: bool, ctx: dict):
    s = d.slide(WHITE)
    d.pic(s, BRAND / "logo-purple.png", X0, 0.5, h=0.36)
    if ctx.get("client_logo"):
        d.pic(s, ctx["client_logo"], X0 + 1.55, 0.48, h=0.4)
    tx, tw = 3.4, SW - 3.4 - X0
    if first:
        d.text(s, tx, 0.28, tw, 0.28, [{"text": f"{num}  ·  {section.get('title_en') or ''}".strip(" ·"), "size": 10.5, "color": MID}], margin=0)
        d.text(s, tx, 0.52, tw, 0.62, [{"text": section.get("title", ""), "size": 28, "bold": True, "color": PURPLE, "line": 1.0}], margin=0)
        if section.get("subtitle"):
            d.text(s, tx, 1.12, tw, 0.36, [{"text": section["subtitle"], "size": 12.5, "color": MUTED}], margin=0)
    else:
        d.text(s, tx, 0.48, tw, 0.5, [{"text": f"{section.get('title', '')}  (تابع)", "size": 18, "bold": True, "color": DEEP}], margin=0)
    d.rect(s, X0, 7.02, CW, 0.012, LINE, radius=0)
    d.text(s, X0, 7.06, 1.0, 0.3, [{"text": f"{d.page:02d}", "size": 10.5, "bold": True, "color": PURPLE, "align": "l"}], margin=0)
    d.text(s, SW - X0 - 7, 7.06, 7, 0.3, [{"text": ctx["footer"], "size": 9.5, "color": MUTED}], margin=0)
    if section.get("type") in ("intro", "conclusion"):
        d.pic(s, BRAND / "mark-purple.png", -0.6, 4.6, h=3.4, alpha=0.05)
    return s


def purple_bg(d: Deck, s, image=None):
    if image:
        d.pic(s, image, 0, 0, w=SW, h=SH)
        d.rect(s, 0, 0, SW, SH, DEEP, radius=0, alpha=0.82)
    d.pic(s, BRAND / "mark-white.png", -0.9, 3.4, h=5.2, alpha=0.08)


def cover_slide(d: Deck, doc: dict, settings: dict, number: str):
    c = doc.get("cover") or {}
    s = d.slide(PURPLE)
    purple_bg(d, s, _local(c.get("image")))
    d.pic(s, BRAND / "logo-white.png", SW - X0 - 2.1, 0.7, h=0.58)
    if _local(c.get("client_logo")):
        box = d.rect(s, X0, 0.62, 2.0, 0.74, WHITE, radius=0.12)
        d.pic(s, _local(c["client_logo"]), X0 + 0.2, 0.72, h=0.54)
    d.text(s, X0, 2.35, CW, 0.4, [{"text": c.get("kicker") or "عرض فني", "size": 15, "color": "E5D6FF"}], margin=0)
    d.text(s, X0, 2.75, CW, 1.4, [{"text": c.get("title", ""), "size": 46, "bold": True, "color": WHITE, "line": 1.0}], margin=0)
    if c.get("subtitle"):
        d.text(s, X0 + 2.5, 4.1, CW - 2.5, 0.9, [{"text": c["subtitle"], "size": 17, "color": "EDE3FF"}], margin=0)
    d.rect(s, X0, 6.0, CW, 0.01, "8C6BC8", radius=0)
    cols = [("مقدم إلى", c.get("client", "")), ("مقدم من", c.get("prepared_by") or settings.get("company_name", "")),
            ("التاريخ", c.get("date", "")), ("النسخة", f"{c.get('version', '1.0')} · {number}")]
    colw = 2.6
    for i, (k, v) in enumerate(cols):
        x = SW - X0 - (i + 1) * colw - i * 0.15
        d.text(s, x, 6.15, colw, 0.3, [{"text": k, "size": 10.5, "color": "CDB8F2"}], margin=0)
        d.text(s, x, 6.43, colw, 0.5, [{"text": v, "size": 12.5, "bold": True, "color": WHITE}], margin=0)
    d.text(s, X0, 6.15, 2.4, 0.3, [{"text": settings.get("website", "www.ahdaf.co"), "size": 12, "bold": True, "color": WHITE, "align": "l"}], margin=0)
    d.text(s, X0, 6.43, 2.4, 0.3, [{"text": settings.get("tagline", "نربط الإبداع بالهدف"), "size": 10.5, "color": "CDB8F2", "align": "l"}], margin=0)


def divider_slide(d: Deck, section: dict, num: str):
    s = d.slide(PURPLE)
    purple_bg(d, s, _local(section.get("image")))
    d.pic(s, BRAND / "logo-white.png", X0, 0.6, h=0.4)
    d.text(s, X0, 2.55, CW - 0.4, 0.45, [{"text": num, "size": 17, "color": "D7C2FF"}], margin=0)
    d.text(s, X0, 3.0, CW - 0.4, 1.2, [{"text": section.get("title", ""), "size": 50, "bold": True, "color": WHITE, "line": 1.0}], margin=0)
    if section.get("subtitle"):
        d.text(s, X0 + 3, 4.25, CW - 3.4, 0.8, [{"text": section["subtitle"], "size": 18, "color": "E8DCFF"}], margin=0)


def closing_slide(d: Deck, doc: dict, settings: dict):
    s = d.slide(PURPLE)
    purple_bg(d, s)
    d.pic(s, BRAND / "logo-white.png", SW / 2 - 1.7, 1.4, h=0.95)
    d.text(s, 1.5, 2.55, SW - 3, 0.4, [{"text": settings.get("tagline", "نربط الإبداع بالهدف"), "size": 16, "color": "E1D0FF", "align": "ctr"}], margin=0)
    d.text(s, 1.5, 3.3, SW - 3, 1.0, [{"text": "شكرًا لكم", "size": 44, "bold": True, "color": WHITE, "align": "ctr"}], margin=0)
    cover = doc.get("cover") or {}
    line = f"نتطلع إلى العمل مع {_plain(cover.get('client'))} على مشروع «{_plain(cover.get('title'))}»."
    d.text(s, 1.5, 4.35, SW - 3, 0.5, [{"text": line, "size": 15, "color": "E8DCFF", "align": "ctr"}], margin=0)
    contacts = "    ".join(x for x in [settings.get("website"), settings.get("email"), settings.get("phone")] if x)
    d.text(s, 1.5, 5.4, SW - 3, 0.4, [{"text": contacts, "size": 13, "color": "F2EAFF", "align": "ctr"}], margin=0)
    if settings.get("address"):
        d.text(s, 1.5, 5.8, SW - 3, 0.4, [{"text": settings["address"], "size": 12, "color": "CDB8F2", "align": "ctr"}], margin=0)


def _local(url):
    """Map an uploaded file URL (/files/<name>) to its path on disk."""
    if not url or not str(url).startswith("/files/"):
        return None
    p = UPLOADS / Path(str(url)).name
    return p if p.exists() else None


# ---------------------------------------------------------------------- block layout
class Flow:
    """Places blocks top-to-bottom on content slides, opening continuation slides as needed."""

    def __init__(self, d: Deck, section: dict, num: str, ctx: dict):
        self.d, self.section, self.num, self.ctx = d, section, num, ctx
        self.s = content_slide(d, section, num, True, ctx)
        self.y = Y0
        self.fresh = True

    def room(self) -> float:
        return YMAX - self.y

    def new(self):
        self.s = content_slide(self.d, self.section, self.num, False, self.ctx)
        self.y = 1.25
        self.fresh = True

    def ensure(self, h: float):
        if h > self.room() and not self.fresh:
            self.new()

    def advance(self, h: float):
        self.y += h + GAP
        self.fresh = False


def place_block(f: Flow, b: dict):
    t = b.get("type")
    d = f.d
    if t in ("paragraph", "lead"):
        size = 18 if t == "lead" else 13.5
        color = DEEP if t == "lead" else INK
        text = b.get("text", "")
        h = th(text, CW, size, 1.45)
        if h > f.room() and f.fresh is False:
            f.new()
        if h > f.room():  # very long paragraph: split by sentences
            parts = re.split(r"(?<=[.!؟?])\s+", text)
            chunk = []
            for part in parts:
                if th(" ".join(chunk + [part]), CW, size, 1.45) > f.room() and chunk:
                    hh = th(" ".join(chunk), CW, size, 1.45)
                    d.text(f.s, X0, f.y, CW, hh, [{"text": " ".join(chunk), "size": size, "color": color, "line": 1.25, "bold": t == "lead"}])
                    f.advance(hh)
                    f.new()
                    chunk = []
                chunk.append(part)
            text = " ".join(chunk)
            h = th(text, CW, size, 1.45)
        d.text(f.s, X0, f.y, CW, h, [{"text": text, "size": size, "color": color, "line": 1.25, "bold": t == "lead"}])
        f.advance(h)
    elif t == "callout":
        h = max(0.6, th(b.get("text"), CW - 2.0, 13.5, 1.45) + 0.3)
        f.ensure(h)
        d.rect(f.s, X0, f.y, CW, h, PURPLE, radius=0.14)
        d.rect(f.s, SW - X0 - 1.35, f.y + 0.16, 1.15, 0.32, WHITE, radius=0.16)
        d.text(f.s, SW - X0 - 1.35, f.y + 0.16, 1.15, 0.32, [{"text": b.get("label", "ملاحظة"), "size": 10.5, "bold": True, "color": PURPLE, "align": "ctr"}], anchor="middle", margin=0)
        d.text(f.s, X0 + 0.25, f.y + 0.12, CW - 1.85, h - 0.2, [{"text": b.get("text", ""), "size": 13.5, "bold": True, "color": WHITE, "line": 1.2}], anchor="middle")
        f.advance(h)
    elif t == "stats":
        items = b.get("items") or []
        n = max(1, min(4, len(items)))
        w = (CW - GAP * (n - 1)) / n
        f.ensure(1.3)
        shades = [PURPLE, "5E12A6", "7330BA", "8A4ACB"]
        for i, it in enumerate(items[:4]):
            x = SW - X0 - (i + 1) * w - i * GAP
            d.rect(f.s, x, f.y, w, 1.25, shades[i % 4], radius=0.14)
            d.text(f.s, x, f.y + 0.12, w, 0.6, [{"text": it.get("value", ""), "size": 28, "bold": True, "color": WHITE, "align": "ctr", "ltr": True}], margin=0)
            d.text(f.s, x + 0.1, f.y + 0.75, w - 0.2, 0.4, [{"text": it.get("label", ""), "size": 11.5, "color": "EFE6FF", "align": "ctr"}], margin=0)
        f.advance(1.25)
    elif t in ("cards", "objectives", "team", "steps"):
        place_cards(f, b)
    elif t == "list":
        place_list(f, b)
    elif t == "services":
        groups = [{"title": g.get("title", ""), "text": "\n".join(f"• {i.get('title', '')}" for i in g.get("items") or [])} for g in b.get("groups") or []]
        place_cards(f, {"type": "cards", "items": groups, "columns": 3, "small": True})
    elif t == "phases":
        place_phases(f, b)
    elif t in ("deliverables", "kpis", "risks", "table"):
        place_table(f, b)
    elif t == "timeline":
        place_gantt(f, b)
    elif t == "image" and _local(b.get("src")):
        h = min(3.6, f.room())
        f.ensure(2.5)
        d.pic(f.s, _local(b["src"]), X0 + CW / 2 - h * 0.9, f.y, h=h)
        f.advance(h)


def place_list(f: Flow, b: dict):
    items = [str(x) for x in b.get("items") or []]
    if b.get("title"):
        f.ensure(0.8)
        f.d.text(f.s, X0, f.y, CW, 0.35, [{"text": b["title"], "size": 14, "bold": True, "color": DEEP}])
        f.advance(0.3)
    two = len(items) > 5 and all(len(x) < 130 for x in items)
    cols = 2 if two else 1
    colw = (CW - 0.4) / cols if two else CW
    i = 0
    while i < len(items):
        heights = [th(x, colw - 0.3, 13, 1.4) + 0.06 for x in items]
        # take as many items as fit in the remaining room across the columns
        best = 0
        for k in range(1, len(items) - i + 1):
            chunk = heights[i:i + k]
            per = math.ceil(k / cols)
            col_h = max(sum(chunk[c * per:(c + 1) * per]) for c in range(cols))
            if col_h <= f.room():
                best = k
            else:
                break
        if best == 0:
            if f.fresh:
                best = 1
            else:
                f.new()
                continue
        chunk = items[i:i + best]
        per = math.ceil(len(chunk) / cols)
        col_h = 0
        for c in range(cols):
            part = chunk[c * per:(c + 1) * per]
            if not part:
                continue
            h = sum(th(x, colw - 0.3, 13, 1.4) + 0.06 for x in part)
            col_h = max(col_h, h)
            x = SW - X0 - (c + 1) * colw - c * 0.4
            f.d.text(f.s, x, f.y, colw, h, [{"text": t, "size": 13, "color": INK, "bullet": True, "space_after": 5, "line": 1.15} for t in part])
        f.advance(col_h)
        i += best
        if i < len(items):
            f.new()


def _card_content(b: dict, it: dict, i: int):
    t = b.get("type")
    if t == "objectives":
        meta = f"يُقاس بـ: {it.get('measure', '')}   |   المستهدف: {it.get('target', '')}" if it.get("measure") else f"المستهدف: {it.get('target', '')}"
        return f"{i + 1}. {it.get('title', '')}", it.get("text", ""), meta
    if t == "team":
        resp = "\n".join(f"• {r}" for r in it.get("responsibilities") or [])
        title = it.get("role", "") + (f"  ×{it.get('count')}" if int(it.get("count") or 1) > 1 else "")
        return title, resp, it.get("role_en", "")
    if t == "steps":
        return f"{i + 1:02d}  {it.get('title', '')}", it.get("text", ""), ""
    return it.get("title", ""), it.get("text", ""), ""


def place_cards(f: Flow, b: dict):
    items = b.get("items") or b.get("members") or []
    t = b.get("type")
    cols = {"objectives": 2, "team": 3, "steps": 4}.get(t) or b.get("columns") or (2 if len(items) <= 4 else 3)
    cols = max(1, min(cols, 4))
    w = (CW - GAP * (cols - 1)) / cols
    tsz, bsz = (13.5, 11.5) if cols >= 3 else (14.5, 12.5)
    if b.get("compact"):
        tsz, bsz = 12, 10
    if b.get("title"):
        f.ensure(1.2)
        f.d.text(f.s, X0, f.y, CW, 0.35, [{"text": b["title"], "size": 14, "bold": True, "color": DEEP}])
        f.advance(0.3)
    rows = [items[i:i + cols] for i in range(0, len(items), cols)]
    base = 0
    for row in rows:
        contents = [_card_content(b, it, base + j) for j, it in enumerate(row)]
        heights = [0.22 + th(c[0], w - 0.36, tsz, 1.3) + th(c[1], w - 0.36, bsz, 1.45) + (th(c[2], w - 0.36, 10.5, 1.3) if c[2] else 0) + 0.18 for c in contents]
        h = max(heights)
        f.ensure(h)
        for j, (ttl, body, meta) in enumerate(contents):
            x = SW - X0 - (j + 1) * w - j * GAP
            f.d.rect(f.s, x, f.y, w, h, LILAC, radius=0.14)
            yy = f.y + 0.14
            th1 = th(ttl, w - 0.36, tsz, 1.3)
            f.d.text(f.s, x + 0.18, yy, w - 0.36, th1, [{"text": ttl, "size": tsz, "bold": True, "color": PURPLE, "line": 1.05}], margin=0)
            yy += th1 + 0.04
            bh = th(body, w - 0.36, bsz, 1.45)
            if body:
                f.d.text(f.s, x + 0.18, yy, w - 0.36, bh, [{"text": line, "size": bsz, "color": INK2, "line": 1.15, "space_after": 2} for line in str(body).split("\n")], margin=0)
            if meta:
                f.d.text(f.s, x + 0.18, f.y + h - th(meta, w - 0.36, 10.5, 1.3) - 0.1, w - 0.36, th(meta, w - 0.36, 10.5, 1.3),
                         [{"text": meta, "size": 10.5, "color": MUTED}], margin=0)
        f.advance(h)
        base += len(row)


def place_phases(f: Flow, b: dict):
    items = b.get("items") or []
    actw, titw = CW - 3.6, 2.9
    for i, p in enumerate(items):
        acts = p.get("activities") or []
        h_acts = sum(th(a, actw - 0.4, 11.5, 1.35) for a in acts)
        h_title = 0.2 + th(p.get("title"), titw, 14, 1.3) + 0.3 + th(p.get("goal"), titw, 10.5, 1.35)
        h = max(h_acts, h_title) + 0.25
        f.ensure(h)
        f.d.rect(f.s, X0, f.y, CW, h, LILAC, radius=0.14)
        nx = SW - X0 - 0.62
        f.d.rect(f.s, nx, f.y + 0.16, 0.46, 0.46, PURPLE, radius=0.1)
        f.d.text(f.s, nx, f.y + 0.16, 0.46, 0.46, [{"text": f"{(b.get('_offset') or 0) + i + 1:02d}", "size": 13, "bold": True, "color": WHITE, "align": "ctr"}], anchor="middle", margin=0)
        tx = nx - 0.15 - titw
        f.d.text(f.s, tx, f.y + 0.12, titw, h - 0.2, [
            {"text": p.get("title", ""), "size": 14, "bold": True, "color": PURPLE, "line": 1.05},
            {"text": p.get("period", ""), "size": 10.5, "color": MID, "space_after": 3},
            {"text": p.get("goal", ""), "size": 10.5, "color": MUTED, "line": 1.15}], margin=0)
        f.d.text(f.s, X0 + 0.2, f.y + 0.12, actw - 0.4, h - 0.2,
                 [{"text": a, "size": 11.5, "color": INK, "bullet": True, "line": 1.1, "space_after": 3} for a in acts], margin=0)
        f.advance(h)


def _lvl(v: str) -> str:
    return "hi" if "مرتفع" in v else "md" if "متوسط" in v else "lo"


def place_table(f: Flow, b: dict):
    t = b.get("type")
    if t == "deliverables":
        header = ["#", "المخرج", "الوصف", "الكمية", "الوحدة", "التوقيت"]
        fr = [0.05, 0.22, 0.37, 0.08, 0.1, 0.18]
        data = [[str(i + 1), r.get("title", ""), r.get("description", ""), "[[يُحدد]]" if r.get("qty") in (None, "") else str(r.get("qty")),
                 r.get("unit", ""), r.get("timing", "")] for i, r in enumerate(b.get("rows") or [])]
        fmt = {0: {"color": MUTED, "align": "ctr"}, 3: {"bold": True, "color": PURPLE, "align": "ctr"}}
        first_bold_col = 1
    elif t == "kpis":
        header = ["المؤشر", "التعريف", "أداة القياس", "الدورية", "المستهدف"]
        fr = [0.2, 0.37, 0.17, 0.1, 0.16]
        data = [[f"{k.get('name', '')}\n{k.get('name_en', '')}".strip(), k.get("definition", ""), k.get("measurement", ""), k.get("frequency", ""), k.get("target", "")] for k in b.get("items") or []]
        fmt = {4: {"bold": True, "color": PURPLE}}
        first_bold_col = 0
    elif t == "risks":
        header = ["الخطر", "الاحتمالية", "التأثير", "آلية المعالجة"]
        fr = [0.3, 0.12, 0.12, 0.46]
        data = [[r.get("risk", ""), r.get("likelihood", ""), r.get("impact", ""), r.get("mitigation", "")] for r in b.get("rows") or []]
        fmt = {1: {"align": "ctr"}, 2: {"align": "ctr"}}
        first_bold_col = 0
    else:
        cols = b.get("columns") or []
        header = [c.get("label", "") for c in cols]
        fr = [1 / max(1, len(cols))] * len(cols)
        data = [[r.get(c["key"], "") for c in cols] for r in b.get("rows") or []]
        fmt = {}
        first_bold_col = 0
        if b.get("title"):
            f.ensure(1.2)
            f.d.text(f.s, X0, f.y, CW, 0.35, [{"text": b["title"], "size": 14, "bold": True, "color": DEEP}])
            f.advance(0.3)
    colw = [CW * x for x in fr]

    def rh(row):
        return max(0.38, max(th(v, cw - 0.2, 11.5, 1.32) + 0.12 for v, cw in zip(row, colw)))

    i = 0
    while i < len(data):
        avail = f.room() - 0.42
        k, used = 0, 0.0
        while i + k < len(data) and used + rh(data[i + k]) <= avail:
            used += rh(data[i + k])
            k += 1
        if k == 0:
            if f.fresh:
                k = 1
            else:
                f.new()
                continue
        chunk = data[i:i + k]
        heights = [rh(r) for r in chunk]
        cell_fmt = dict(fmt)
        if t == "risks":
            cell_fmt = {}
        h = f.d.table(f.s, X0, f.y, CW, colw, header, [r for r in chunk], heights, first_bold=True, cell_fmt=cell_fmt)
        if t == "risks":  # colour the level cells
            tbl = f.s.shapes[-1].table
            for ri, r in enumerate(chunk, start=1):
                for ci in (1, 2):
                    bg, fg = LVL[_lvl(r[ci])]
                    f.d._cell(tbl.cell(ri, ci), r[ci], 11, True, fg, bg, align="ctr")
        if first_bold_col == 1:
            tbl = f.s.shapes[-1].table
            for ri, r in enumerate(chunk, start=1):
                f.d._cell(tbl.cell(ri, 1), r[1], 11.5, True, DEEP, "FBF9FE" if ri % 2 == 0 else WHITE)
                f.d._cell(tbl.cell(ri, 0), r[0], 11, False, MUTED, "FBF9FE" if ri % 2 == 0 else WHITE, align="ctr")
        f.advance(h)
        i += k
        if i < len(data):
            f.new()


def place_gantt(f: Flow, b: dict):
    from datetime import date, timedelta
    rows = b.get("rows") or []
    total = max(1, int(b.get("total_days") or 1))
    start = None
    if b.get("start_date"):
        try:
            start = date.fromisoformat(str(b["start_date"])[:10])
        except ValueError:
            start = None
    months = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]

    def dt(off):
        if start:
            d = start + timedelta(days=off)
            return f"{d.day} {months[d.month - 1]}"
        return f"أسبوع {off // 7 + 1}"

    lw, dw, ow = 3.6, 2.3, 1.6
    trackw = CW - lw - dw - ow
    track_x = X0
    rowh = 0.44
    i = 0
    while i < len(rows):
        avail = f.room() - 0.42
        k = max(1, min(len(rows) - i, int(avail // rowh)))
        if avail < rowh * 2 and not f.fresh:
            f.new()
            continue
        chunk = rows[i:i + k]
        y = f.y
        # header
        f.d.rect(f.s, X0, y, CW, 0.4, PURPLE, radius=0.08)
        hx = SW - X0
        for label, w in (("المرحلة / النشاط", lw), ("البداية – النهاية", dw), ("المسؤول", ow)):
            f.d.text(f.s, hx - w, y, w, 0.4, [{"text": label, "size": 10.5, "color": WHITE}], anchor="middle", margin=0.08)
            hx -= w
        step = 7 if total <= 70 else 30.4
        n = int(math.ceil(total / step))
        for t_i in range(n):
            at = t_i * step
            x = track_x + trackw - (at + step / 2) / total * trackw
            label = (months[(start + timedelta(days=int(at))).month - 1] if start else f"ش{t_i + 1}") if step > 7 else f"أ{t_i + 1}"
            if x - 0.3 >= track_x:
                f.d.text(f.s, x - 0.35, y, 0.7, 0.4, [{"text": label, "size": 9, "color": "EFE6FF", "align": "ctr"}], anchor="middle", margin=0)
        y += 0.4
        last_phase = None
        for j, r in enumerate(chunk):
            bg = "FBF9FE" if j % 2 else WHITE
            f.d.rect(f.s, X0, y, CW, rowh, bg, radius=0)
            f.d.rect(f.s, X0, y + rowh - 0.01, CW, 0.01, LINE, radius=0)
            label = r.get("activity", "")
            paras = []
            if r.get("phase") != last_phase:
                paras.append({"text": r.get("phase", ""), "size": 9.5, "bold": True, "color": PURPLE, "space_after": 0, "line": 1.0})
            last_phase = r.get("phase")
            paras.append({"text": label, "size": 10.5, "color": INK, "space_after": 0, "line": 1.0})
            f.d.text(f.s, SW - X0 - lw, y, lw, rowh, paras, anchor="middle", margin=0.08)
            off, days = int(r.get("offset") or 0), max(1, int(r.get("days") or 1))
            f.d.text(f.s, SW - X0 - lw - dw, y, dw, rowh, [{"text": f"{dt(off)} – {dt(off + days - 1)}", "size": 10, "color": MUTED}], anchor="middle", margin=0.08)
            f.d.text(f.s, SW - X0 - lw - dw - ow, y, ow, rowh, [{"text": r.get("owner", ""), "size": 10, "color": DEEP}], anchor="middle", margin=0.08)
            bx = track_x + trackw - (off + days) / total * trackw
            bw = max(0.06, days / total * trackw)
            f.d.rect(f.s, max(track_x, bx), y + rowh / 2 - 0.08, min(bw, trackw), 0.16, PURPLE if j % 2 == 0 else ACCENT, radius=0.08)
            y += rowh
        used = 0.4 + rowh * len(chunk)
        last = i + k >= len(rows)
        ms = [m for m in (b.get("milestones") or []) if 0 <= int(m.get("offset") or 0) <= total]
        if last and ms:
            f.d.rect(f.s, X0, y, CW, 0.38, LILAC, radius=0)
            f.d.text(f.s, SW - X0 - lw, y, lw, 0.38, [{"text": "المحطات الرئيسية", "size": 10.5, "bold": True, "color": PURPLE}], anchor="middle", margin=0.08)
            for m in ms:
                cx = track_x + trackw - int(m["offset"]) / total * trackw
                dm = f.d.rect(f.s, cx - 0.07, y + 0.12, 0.14, 0.14, PURPLE, radius=0, shape=MSO_SHAPE.DIAMOND)
            used += 0.38
            legend = "   ◆ ".join(f"{_plain(m['label'])}" + (f" ({dt(int(m['offset']))})" if start else "") for m in ms)
            f.d.text(f.s, X0, y + 0.44, CW, 0.5, [{"text": "◆ " + legend, "size": 9.5, "color": MUTED}], margin=0)
            used += 0.5
        f.advance(used)
        i += k
        if i < len(rows):
            f.new()


# ---------------------------------------------------------------------- public
def build_pptx(proposal: dict, settings: dict) -> bytes:
    doc = proposal.get("doc") or {}
    d = Deck(settings.get("pptx_font") or "Saudi Serif VF")
    sections = [s for s in doc.get("sections", []) if s.get("visible", True)]
    ctx = {"footer": f"عرض فني — {_plain((doc.get('cover') or {}).get('title'))}",
           "client_logo": _local((doc.get("cover") or {}).get("client_logo"))}
    cover_slide(d, doc, settings, proposal.get("number", ""))

    if doc.get("show_toc", True) and len(sections) > 2:
        toc = {"id": "__toc", "type": "toc", "title": "محتويات العرض", "title_en": "Contents"}
        f = Flow(d, toc, "", ctx)
        place_cards(f, {"type": "cards", "columns": 4, "compact": True, "items": [{"title": f"{i + 1:02d}  {s.get('title', '')}", "text": s.get("subtitle", "")} for i, s in enumerate(sections)]})

    for i, s in enumerate(sections):
        num = f"{i + 1:02d}"
        if s.get("divider"):
            divider_slide(d, s, num)
        f = Flow(d, s, num, ctx)
        for b in s.get("blocks", []):
            place_block(f, b)
        if s.get("image") and not s.get("divider"):
            place_block(f, {"type": "image", "src": s["image"]})
    closing_slide(d, doc, settings)

    # document properties
    cp = d.prs.core_properties
    cp.title = _plain((doc.get("cover") or {}).get("title"))
    cp.author = settings.get("company_name", "أهداف")
    cp.subject = f"عرض فني — {_plain((doc.get('cover') or {}).get('client'))}"
    buf = io.BytesIO()
    d.prs.save(buf)
    return buf.getvalue()
